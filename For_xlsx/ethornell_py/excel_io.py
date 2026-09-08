"""Excel workbook read/write matching VNTextPatch ExcelScript.cs and ExcelScriptCollection.cs.

Columns:
  A (1): OriginalCharacter
  B (2): OriginalLine
  C (3): TranslatedCharacter
  D (4): TranslatedLine (TL)
  E (5): CheckedLine (TLC)
  F (6): EditedLine (Edit)
  G (7): Notes
"""

import os
import re
from typing import List, Dict, Optional, Tuple, Iterable
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from .disassembler import ScriptString, ScriptStringType
from .names_xml import CharacterNames

EMPTY_TEXT_MARKER = "(empty)"


def quote_name(name: str) -> str:
    if "/" not in name and '"' not in name:
        return name
    name = name.replace('\\', '\\\\').replace('"', '\\"')
    return f'"{name}"'


def unquote_name(name: str) -> str:
    if not (name.startswith('"') and name.endswith('"')):
        return name
    name = name[1:-1]
    name = re.sub(r'\\(.)', r'\1', name)
    return name


def join_names(names: Iterable[str]) -> str:
    return "/".join(quote_name(n) for n in names)


def split_names(names: str) -> List[str]:
    matches = re.findall(r'(?:"(?:\\.|[^"])+"|[^/]+)', names)
    return [unquote_name(m) for m in matches]


class ExcelHandler:
    @staticmethod
    def get_row_text(row_values: Dict[int, Optional[str]]) -> str:
        """Priority: Edit (if not '.') -> TLC (if not '.') -> TL -> Original.

        If a row exists in the spreadsheet, it corresponds to a scenario line.
        If all text columns are empty or None, return empty string '' so that
        the line count remains synchronized with the binary script.
        """
        orig_text = row_values.get(2)
        tl_text = row_values.get(4)
        tlc_text = row_values.get(5)
        edit_text = row_values.get(6)

        def null_if(val, check):
            return None if val == check else val

        def null_if_none_or_empty(val):
            if val is None:
                return None
            s = str(val)
            return None if len(s) == 0 else s

        tl_text = null_if_none_or_empty(tl_text)
        tlc_text = null_if_none_or_empty(tlc_text)
        edit_text = null_if_none_or_empty(edit_text)

        text = null_if(edit_text, ".") or null_if(tlc_text, ".") or tl_text
        if text is None:
            if orig_text is None:
                return ""
            text = str(orig_text)

        if text == EMPTY_TEXT_MARKER:
            return ""
        return text

    @classmethod
    def read_strings_from_sheet(cls, worksheet) -> List[ScriptString]:
        results: List[ScriptString] = []
        max_row = worksheet.max_row
        if max_row < 2:
            return results

        for row in worksheet.iter_rows(min_row=2, max_row=max_row, values_only=True):
            row_dict = {col_idx + 1: val for col_idx, val in enumerate(row)}

            tl_char = row_dict.get(3)
            orig_char = row_dict.get(1)

            def null_if_none_or_empty(val):
                if val is None:
                    return None
                s = str(val)
                return None if len(s) == 0 else s

            char_names_str = null_if_none_or_empty(tl_char) or null_if_none_or_empty(orig_char)

            if char_names_str:
                names = split_names(char_names_str)
                for name in names:
                    results.append(ScriptString(name, ScriptStringType.CharacterName))

            text = cls.get_row_text(row_dict)
            text = re.sub(r'(?<!\r)\n', '\r\n', text)
            results.append(ScriptString(text, ScriptStringType.Message))

        return results

    @classmethod
    def read_workbook(cls, file_path: str) -> Dict[str, List[ScriptString]]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheets_data = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheets_data[sheet_name] = cls.read_strings_from_sheet(ws)
        return sheets_data

    @classmethod
    def write_workbook(
        cls,
        file_path: str,
        scripts_data: Dict[str, List[ScriptString]],
        char_names_db: Optional[CharacterNames] = None,
        template_path: Optional[str] = None
    ) -> None:
        if template_path and os.path.exists(template_path):
            wb = openpyxl.load_workbook(template_path)
            template_ws = wb.active
        else:
            wb = openpyxl.Workbook()
            template_ws = wb.active
            template_ws.title = "Script"
            headers = [None, "Original", None, "TL", "TLC", "Edit", "Notes"]
            for col_idx, header in enumerate(headers, 1):
                if header:
                    template_ws.cell(row=1, column=col_idx, value=header)

        first_sheet = True
        for sheet_name, strings in scripts_data.items():
            if first_sheet:
                ws = wb.active
                ws.title = sheet_name
                first_sheet = False
            else:
                ws = wb.copy_worksheet(template_ws)
                ws.title = sheet_name

            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row)

            row_num = 2
            pending_names: List[str] = []
            for item in strings:
                if item.type == ScriptStringType.CharacterName:
                    pending_names.append(item.text)
                else:
                    orig_char_str = join_names(pending_names) if pending_names else None
                    msg_text = item.text if item.text else EMPTY_TEXT_MARKER

                    if orig_char_str:
                        ws.cell(row=row_num, column=1, value=orig_char_str)

                    ws.cell(row=row_num, column=2, value=msg_text)

                    if orig_char_str and char_names_db:
                        trans_names = [char_names_db.get_translation(n) for n in pending_names]
                        ws.cell(row=row_num, column=3, value=join_names(trans_names))

                    pending_names.clear()
                    row_num += 1

        wb.save(file_path)
